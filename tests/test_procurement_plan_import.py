import importlib.util
from pathlib import Path

import openpyxl
import pytest


MODULE_PATH = Path(__file__).parents[1] / 'app' / 'procurement' / 'plan_import.py'
SPEC = importlib.util.spec_from_file_location('procurement_plan_import', MODULE_PATH)
PLAN_IMPORT = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(PLAN_IMPORT)


def _workbook(path, rows, header_row=2):
    workbook = openpyxl.Workbook()
    summary = workbook.active
    summary.title = 'งบลงทุน'
    summary['A1'] = 'รายการงบลงทุน ประจำปีงบประมาณ พ.ศ. 2570'
    sheet = workbook.create_sheet(PLAN_IMPORT.DATA_SHEET)
    headers = [
        'Long Text',
        'ผลผลิต/โครงการ/รายงาน (Functional Area)',
        'หน่วยงานที่รับผิดชอบ\n(Cost center)',
        'วิธีการจัดซื้อจัดจ้าง (แผน)',
        'จำนวนเงิน',
        'Fund',
        'วันที่ประกาศ/วันที่อนุมัติ (แผน)',
        'วันที่ลงนาม\nในสัญญา (แผน)',
        'วันที่ตรวจรับงวด 1 (แผน)',
        'วันที่ตรวจรับงวด 2 (แผน)',
        'External order no.',
    ]
    for index in range(header_row - 1):
        sheet.append(['รายละเอียดงบประมาณ' if index == 0 else None])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


@pytest.mark.parametrize('header_row', [1, 3, 5])
def test_detects_header_after_title_and_blank_rows(tmp_path, header_row):
    path = tmp_path / 'plans.xlsx'
    _workbook(path, [
        ['เงินแผ่นดิน (20101002)'],
        ['รายการทดสอบ', '0150001 Test', 'C0401000',
         'วิธีเฉพาะเจาะจง', 1000, '20101002'],
    ], header_row=header_row)
    # Section labels in the real workbook are outside the Long Text column.
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[PLAN_IMPORT.DATA_SHEET]
    sheet.cell(header_row + 1, 1).value = None
    sheet.cell(header_row + 1, 12).value = 'เงินแผ่นดิน (20101002)'
    workbook.save(path)

    _, rows = PLAN_IMPORT.parse_procurement_plan_workbook(path, 2570)

    assert len(rows) == 1
    assert rows[0]['item'] == 'รายการทดสอบ'
    assert rows[0]['row_number'] == header_row + 2


def test_reports_excel_row_number_after_detected_header(tmp_path):
    path = tmp_path / 'plans.xlsx'
    _workbook(path, [['รายการทดสอบ', '0150001 Test', 'C0401000',
                      'วิธีเฉพาะเจาะจง', 'ไม่ใช่ตัวเลข', '20101002']], header_row=3)

    with pytest.raises(PLAN_IMPORT.ProcurementPlanImportError, match='แถวที่ 4'):
        PLAN_IMPORT.parse_procurement_plan_workbook(path, 2570)


def test_rejects_missing_required_column(tmp_path):
    path = tmp_path / 'plans.xlsx'
    _workbook(path, [], header_row=3)
    workbook = openpyxl.load_workbook(path)
    workbook[PLAN_IMPORT.DATA_SHEET].cell(3, 6).value = None
    workbook.save(path)

    with pytest.raises(PLAN_IMPORT.ProcurementPlanImportError, match='ไม่พบคอลัมน์ที่จำเป็น: Fund'):
        PLAN_IMPORT.parse_procurement_plan_workbook(path, 2570)


def test_parse_standard_workbook(tmp_path):
    path = tmp_path / 'plans.xlsx'
    _workbook(path, [[
        'กล้องจุลทรรศน์ 1 ชุด',
        '0150001 วิทยาศาสตร์สุขภาพLS',
        'C0401301',
        'วิธีประกวดราคาอิเล็กทรอนิกส์ (e-bidding)',
        2490000,
        20101002,
        '31/10/2026',
        '31/01/2027',
        '31/03/2027',
        '31/05/2027',
        'IO-123',
    ]])

    with path.open('rb') as workbook:
        fiscal_year, rows = PLAN_IMPORT.parse_procurement_plan_workbook(workbook, 2571)

    assert fiscal_year == 2571
    assert len(rows) == 1
    assert rows[0]['product_code'] == '0150001'
    assert rows[0]['product_name'] == 'วิทยาศาสตร์สุขภาพLS'
    assert rows[0]['cost_center'] == 'C0401301'
    assert rows[0]['funding_source'] == '20101002'
    assert rows[0]['fund_code'] == 'IO-123'
    assert 'principle_approval_date' not in rows[0]
    assert 'contract_signed_date' not in rows[0]
    assert 'inspection_date' not in rows[0]


def test_rejects_invalid_amount_without_returning_partial_rows(tmp_path):
    path = tmp_path / 'plans.xlsx'
    _workbook(path, [[
        'รายการทดสอบ', '0150001 Test', 'C0401000',
        'วิธีเฉพาะเจาะจง', 'ไม่ใช่ตัวเลข', '20101002',
        None, None, None, None, None,
    ]])

    with path.open('rb') as workbook, pytest.raises(PLAN_IMPORT.ProcurementPlanImportError) as error:
        PLAN_IMPORT.parse_procurement_plan_workbook(workbook, 2570)

    assert 'แถวที่ 3' in str(error.value)
    assert 'จำนวนเงิน' in str(error.value)


@pytest.mark.parametrize('source, expected', [
    ('วิธีประกวดราคาอิเล็กทรอนิกส์ (e-bidding)', 'ประกาศเชิญชวนทั่วไป(E-Bidding)'),
    ('วิธีคัดเลือก', 'วิธีคัดเลือก'),
    ('วิธีเฉพาะเจาะจง', 'วิธีเฉพาะเจาะจง'),
])
def test_normalises_procurement_method(source, expected):
    assert PLAN_IMPORT.normalise_procurement_method(source) == expected
